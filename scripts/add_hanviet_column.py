#!/usr/bin/env python3
"""Add a Hán Việt column to the HSK vocabulary workbook.

Readings come from https://github.com/ph0ngp/hanviet-pinyin-words
(traditional characters + numbered pinyin). Words with no complete
reading are left blank.
"""

from __future__ import annotations

import argparse
import json
import re
import sys
import urllib.request
from pathlib import Path

from openpyxl import load_workbook

ROOT = Path(__file__).resolve().parents[1]
DEFAULT_XLSX = ROOT / "resources" / "vocabulary_hsk1-3_3.0_and_hsk1-6_2.0.xlsx"
CACHE_DIR = ROOT / "resources" / ".hanviet-cache"
HANVIET_URL = (
    "https://raw.githubusercontent.com/ph0ngp/hanviet-pinyin-words/main/src/hanvietData.js"
)
ST_URL = (
    "https://raw.githubusercontent.com/BYVoid/OpenCC/master/data/dictionary/STCharacters.txt"
)
COLUMN_TITLE = "Hán Việt"

_TONE_LOWER = {
    "ā": ("a", "1"),
    "á": ("a", "2"),
    "ǎ": ("a", "3"),
    "à": ("a", "4"),
    "ē": ("e", "1"),
    "é": ("e", "2"),
    "ě": ("e", "3"),
    "è": ("e", "4"),
    "ī": ("i", "1"),
    "í": ("i", "2"),
    "ǐ": ("i", "3"),
    "ì": ("i", "4"),
    "ō": ("o", "1"),
    "ó": ("o", "2"),
    "ǒ": ("o", "3"),
    "ò": ("o", "4"),
    "ū": ("u", "1"),
    "ú": ("u", "2"),
    "ǔ": ("u", "3"),
    "ù": ("u", "4"),
    "ǖ": ("ü", "1"),
    "ǘ": ("ü", "2"),
    "ǚ": ("ü", "3"),
    "ǜ": ("ü", "4"),
    "ń": ("n", "2"),
    "ň": ("n", "3"),
    "ǹ": ("n", "4"),
    "ḿ": ("m", "2"),
    "m̄": ("m", "1"),
    "m̀": ("m", "4"),
    "ê̄": ("ê", "1"),
    "ế": ("ê", "2"),
    "ê̌": ("ê", "3"),
    "ề": ("ê", "4"),
}
TONE_TO_BASE = dict(_TONE_LOWER)
TONE_TO_BASE.update({key.upper(): value for key, value in _TONE_LOWER.items()})

# OpenCC misses a few simplified forms used in HSK lists.
EXTRA_S2T = {
    "着": ["著"],
    "痹": ["痺"],
    "俩": ["倆"],
}

# Syllables that appear in HSK pinyin but not as keys in hanvietData
# (those entries often use "*" for every reading).
STANDARD_SYLLABLES = """
a o e ê er ai ei ao ou an en ang eng ong
ba bo bai bei bao ban ben bang beng bi bie biao bian bin bing bu
pa po pai pei pao pou pan pen pang peng pi pie piao pian pin ping pu
ma mo me mai mei mao mou man men mang meng mi mie miao miu mian min ming mu
fa fo fei fou fan fen fang feng fu
da de dai dei dao dou dan den dang deng dong di dia die diao diu dian ding du duo dui duan dun
ta te tai tei tao tou tan tang teng tong ti tie tiao tian ting tu tuo tui tuan tun
na ne nai nei nao nou nan nen nang neng nong ni nie niao niu nian nin niang ning nu nuo nuan nun nü nüe nv nve
la le lai lei lao lou lan lang leng long li lia lie liao liu lian lin liang ling lu luo lun luan lü lüe lv lve
ga ge gai gei gao gou gan gen gang geng gong gua guo guai gui guan gun guang gu
ka ke kai kei kao kou kan ken kang keng kong kua kuo kuai kui kuan kun kuang ku
ha he hai hei hao hou han hen hang heng hong hua huo huai hui huan hun huang hu
ji jia jie jiao jiu jian jin jiang jing jiong ju juan jue jun
qi qia qie qiao qiu qian qin qiang qing qiong qu quan que qun
xi xia xie xiao xiu xian xin xiang xing xiong xu xuan xue xun
zhi zha zhe zhao zhou zhan zhen zhang zheng zhong zhu zhua zhuo zhuai zhui zhuan zhun zhuang
chi cha che chao chou chan chen chang cheng chong chu chua chuo chuai chui chuan chun chuang
shi sha she shei shao shou shan shen shang sheng shu shua shuo shuai shui shuan shun shuang
ri re rao rou ran ren rang reng rong ru rua ruo rui ruan run
zi za ze zai zei zao zou zan zen zang zeng zong zu zuo zui zuan zun
ci ca ce cai cao cou can cen cang ceng cong cu cuo cui cuan cun
si sa se sai sao sou san sen sang seng song su suo sui suan sun
ya yo ye yao you yan yin yang ying yong yu yuan yue yun yi
wa wo wai wei wan wen wang weng wu
ng hm hng
""".split()

APOSTROPHES = str.maketrans({"’": "'", "‘": "'", "ʼ": "'", "ˈ": "'", "´": "'"})
PAREN_RE = re.compile(r"[\(（][^）)]*[\)）]")
HANZI_COL_NAMES = {"Từ mới", "TỪ MỚI"}
PINYIN_COL_NAMES = {"Phiên âm", "PHIÊN ÂM"}
HANVIET_COL_NAMES = {"Hán Việt", "HÁN VIỆT", "Han Viet", "hanviet"}


def cell_str(value) -> str | None:
    if value is None:
        return None
    text = str(value).strip()
    return text or None


def download(url: str, dest: Path) -> None:
    dest.parent.mkdir(parents=True, exist_ok=True)
    with urllib.request.urlopen(url) as resp:
        dest.write_bytes(resp.read())


def load_hanviet_data(path: Path) -> dict:
    text = path.read_text(encoding="utf-8").strip()
    prefix = "export const hanvietData = "
    if text.startswith(prefix):
        text = text[len(prefix) :]
    if text.endswith(";"):
        text = text[:-1]
    return json.loads(text)


def load_s2t(path: Path) -> dict[str, list[str]]:
    mapping: dict[str, list[str]] = {}
    for line in path.read_text(encoding="utf-8").splitlines():
        line = line.strip()
        if not line or line.startswith("#"):
            continue
        simplified, traditional = line.split("\t", 1)
        mapping[simplified] = traditional.split()
    return mapping


def traditional_candidates(ch: str, s2t: dict[str, list[str]]) -> list[str]:
    seen: list[str] = []
    for cand in EXTRA_S2T.get(ch, []) + s2t.get(ch, []) + [ch]:
        if cand not in seen:
            seen.append(cand)
    return seen


PINYIN_BASE_ALIASES = {
    "you": ["yo"],
    "yo": ["you"],
}


def normalize_lookup_pinyin(syllable: str) -> list[str]:
    """Return pinyin spellings to try against hanvietData keys."""
    s = syllable.lower().replace("v", "ü")
    variants = [s]
    if "ü" in s:
        variants.append(s.replace("ü", "u:"))
    if "u:" in s:
        variants.append(s.replace("u:", "ü"))
    base, tone = split_tone(s)
    for alias in PINYIN_BASE_ALIASES.get(base, []):
        variants.append(alias + tone)
    out: list[str] = []
    for item in variants:
        if item not in out:
            out.append(item)
    return out


def split_tone(syllable: str) -> tuple[str, str]:
    if syllable and syllable[-1] in "12345":
        return syllable[:-1], syllable[-1]
    return syllable, "5"


def readings_for_pinyin(char_data: dict, syllable: str) -> list[str] | None:
    keys = list(char_data.keys())
    if keys == ["*"]:
        values = [item for item in char_data["*"] if item]
        return values or None

    variants = normalize_lookup_pinyin(syllable)
    for variant in variants:
        values = [item for item in char_data.get(variant, []) if item]
        if values:
            return values

    bases = []
    for variant in variants:
        base, _tone = split_tone(variant)
        if base not in bases:
            bases.append(base)

    seen: list[str] = []
    for base in bases:
        for tone in "51234":
            key = base + tone
            for item in char_data.get(key, []):
                if item and item not in seen:
                    seen.append(item)
    if seen:
        return seen
    fallback: list[str] = []
    for values in char_data.values():
        for item in values:
            if item and item not in fallback:
                fallback.append(item)
    # Only fill when the character has a single Hán Việt reading.
    if len(fallback) == 1:
        return fallback
    return None


def lookup_char(
    ch: str,
    syllable: str,
    hanviet: dict,
    s2t: dict[str, list[str]],
) -> list[str] | None:
    if len(ch) == 1 and 32 <= ord(ch) <= 126:
        return [ch]
    if ch in {"，", ","}:
        return [","]
    if ch in {"。", "."}:
        return ["."]

    for cand in traditional_candidates(ch, s2t):
        char_data = hanviet.get(cand)
        if not char_data:
            continue
        readings = readings_for_pinyin(char_data, syllable)
        if readings:
            return readings
    return None


def toneless_and_tones(chunk: str) -> tuple[str, list[str | None]]:
    letters: list[str] = []
    tones: list[str | None] = []
    i = 0
    while i < len(chunk):
        pair = None
        if i + 1 < len(chunk):
            pair = TONE_TO_BASE.get(chunk[i : i + 2])
        mapped = pair or TONE_TO_BASE.get(chunk[i])
        if mapped:
            base, tone = mapped
            letters.append(base)
            tones.append(tone)
            i += 2 if pair else 1
            continue
        ch = chunk[i]
        if ch in "12345" and letters:
            tones[-1] = ch
            i += 1
            continue
        if ch == "v":
            letters.append("ü")
        else:
            letters.append(ch.lower())
        tones.append(None)
        i += 1
    return "".join(letters), tones


def _syllable_with_tone(letters: str, tones: list[str | None], start: int, size: int) -> str:
    tone = "5"
    for mark in tones[start : start + size]:
        if mark:
            tone = mark
    return letters[start : start + size] + tone


def segment_chunk(letters: str, tones: list[str | None], valid: set[str]) -> list[list[str]]:
    if letters in {"n", "m"}:
        valid = valid | {letters}
    n = len(letters)
    max_len = max((len(item) for item in valid), default=6)
    memo: dict[int, list[list[str]]] = {}

    def dfs(i: int) -> list[list[str]]:
        if i == n:
            return [[]]
        if i in memo:
            return memo[i]
        found: list[list[str]] = []
        for size in range(min(max_len + 1, n - i), 0, -1):
            piece = letters[i : i + size]
            is_erhua = (
                i + size == n
                and piece.endswith("r")
                and piece != "er"
                and piece[:-1] in valid
            )
            if piece not in valid and not is_erhua:
                continue
            syl = _syllable_with_tone(letters, tones, i, size)
            for rest in dfs(i + size):
                found.append([syl] + rest)
        memo[i] = found
        return found

    return dfs(0)


def pinyin_segmentations(pinyin: str, valid: set[str]) -> list[list[str]]:
    text = pinyin.translate(APOSTROPHES).strip()
    text = re.split(r"[/／]", text, maxsplit=1)[0]
    text = PAREN_RE.sub(" ", text)
    text = " ".join(text.replace("-", " ").split())
    if not text:
        return []

    chunk_options: list[list[list[str]]] = []
    for raw in re.split(r"[\s']+", text):
        if not raw:
            continue
        letters, tones = toneless_and_tones(raw)
        letters = letters.replace("u:", "ü")
        parts = segment_chunk(letters, tones, valid)
        if not parts:
            return []
        chunk_options.append(parts)

    combined: list[list[str]] = [[]]
    for options in chunk_options:
        next_combined: list[list[str]] = []
        for prefix in combined:
            for option in options:
                next_combined.append(prefix + option)
        combined = next_combined
    return combined


def split_erhua(hanzi: list[str], syllables: list[str]) -> list[str]:
    if not hanzi or hanzi[-1] not in {"儿", "兒"}:
        return syllables
    if len(syllables) == len(hanzi):
        return syllables
    if len(syllables) + 1 != len(hanzi) or not syllables:
        return syllables
    last = syllables[-1]
    base, tone = split_tone(last)
    if base.endswith("r") and base not in {"er", "r"}:
        return syllables[:-1] + [base[:-1] + tone, "r5"]
    return syllables


def word_hanzi(raw: str) -> str:
    text = PAREN_RE.sub("", raw)
    return "".join(text.split())


def format_hanviet(parts: list[list[str]]) -> str:
    if len(parts) == 1:
        return ", ".join(parts[0])
    return " ".join(" | ".join(item) for item in parts)


def hanviet_for_word(
    hanzi: str,
    pinyin: str,
    hanviet: dict,
    s2t: dict[str, list[str]],
    valid: set[str],
) -> str:
    chars = list(word_hanzi(hanzi))
    if not chars:
        return ""
    for syllables in pinyin_segmentations(pinyin, valid):
        syllables = split_erhua(chars, syllables)
        if len(chars) != len(syllables):
            continue
        parts: list[list[str]] = []
        ok = True
        for ch, syllable in zip(chars, syllables):
            readings = lookup_char(ch, syllable, hanviet, s2t)
            if not readings:
                ok = False
                break
            parts.append(readings)
        if ok:
            return format_hanviet(parts)
    return ""


def build_valid_syllables(hanviet: dict) -> set[str]:
    valid = set(STANDARD_SYLLABLES)
    for char_data in hanviet.values():
        for key in char_data:
            if key == "*":
                continue
            base, _tone = split_tone(key.lower().replace("u:", "ü").replace("v", "ü"))
            if base and base not in {"n", "m", "r"}:
                valid.add(base)
    valid.update({"ng", "hm", "hng"})
    return valid


def find_header_row(ws) -> int | None:
    for row in range(1, min(ws.max_row, 8) + 1):
        values = {cell_str(cell.value) for cell in ws[row]}
        if values & HANZI_COL_NAMES and values & PINYIN_COL_NAMES:
            return row
    return None


def col_index(ws, header_row: int, names: set[str]) -> int | None:
    for cell in ws[header_row]:
        if cell_str(cell.value) in names:
            return cell.column
    return None


def last_used_col(ws, header_row: int) -> int:
    last = 1
    for cell in ws[header_row]:
        if cell.value not in (None, ""):
            last = cell.column
    return last


def process_workbook(xlsx: Path, hanviet: dict, s2t: dict[str, list[str]], apply: bool) -> int:
    valid = build_valid_syllables(hanviet)
    wb = load_workbook(xlsx)
    filled = 0
    empty = 0
    words = 0

    for sheet in wb.sheetnames:
        ws = wb[sheet]
        header_row = find_header_row(ws)
        if header_row is None:
            print(f"skip {sheet}: no vocabulary header", file=sys.stderr)
            continue
        hanzi_col = col_index(ws, header_row, HANZI_COL_NAMES)
        pinyin_col = col_index(ws, header_row, PINYIN_COL_NAMES)
        if not hanzi_col or not pinyin_col:
            print(f"skip {sheet}: missing hanzi/pinyin column", file=sys.stderr)
            continue

        hv_col = col_index(ws, header_row, HANVIET_COL_NAMES)
        if hv_col is None:
            hv_col = last_used_col(ws, header_row) + 1
            ws.cell(header_row, hv_col, COLUMN_TITLE)

        sheet_filled = 0
        sheet_words = 0
        for row in range(header_row + 1, ws.max_row + 1):
            hanzi = cell_str(ws.cell(row, hanzi_col).value)
            pinyin = cell_str(ws.cell(row, pinyin_col).value)
            if not hanzi:
                continue
            sheet_words += 1
            words += 1
            value = hanviet_for_word(hanzi, pinyin or "", hanviet, s2t, valid)
            if value:
                sheet_filled += 1
                filled += 1
            else:
                empty += 1
            if apply:
                ws.cell(row, hv_col, value)

        print(f"{sheet}: {sheet_filled}/{sheet_words} filled")

    if apply:
        wb.save(xlsx)
        print(f"Wrote {xlsx}")
    else:
        print("Dry run; workbook not modified")
    print(f"Total {filled}/{words} filled, {empty} left empty")
    return 0


def main() -> int:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("xlsx", nargs="?", type=Path, default=DEFAULT_XLSX)
    parser.add_argument("--dry-run", action="store_true")
    parser.add_argument("--refresh-cache", action="store_true")
    args = parser.parse_args()
    xlsx = args.xlsx.resolve()
    if not xlsx.exists():
        print(f"Missing workbook: {xlsx}", file=sys.stderr)
        return 1

    CACHE_DIR.mkdir(parents=True, exist_ok=True)
    data_path = CACHE_DIR / "hanvietData.js"
    st_path = CACHE_DIR / "STCharacters.txt"
    if args.refresh_cache or not data_path.exists():
        print(f"Downloading {HANVIET_URL}")
        download(HANVIET_URL, data_path)
    if args.refresh_cache or not st_path.exists():
        print(f"Downloading {ST_URL}")
        download(ST_URL, st_path)

    hanviet = load_hanviet_data(data_path)
    s2t = load_s2t(st_path)
    return process_workbook(xlsx, hanviet, s2t, apply=not args.dry_run)


if __name__ == "__main__":
    raise SystemExit(main())
